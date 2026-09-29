"""Local annotation projects. Model proposals never become expert labels implicitly."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib,json,uuid,shutil,threading,math,os
import numpy as np
import pydicom
from PIL import Image,ImageOps
from fastapi import APIRouter,HTTPException,Request
from fastapi.responses import FileResponse,Response
from .imaging import read_dicom,geometry
from .specialists import TASKS
from .task_eligibility import applicable_tasks,task_blocks,export_admission

ROOT=Path(__file__).resolve().parents[1]
PROJECTS=Path(os.environ.get('DXA_ANNOTATIONS',str(ROOT/'annotation_projects')))
router=APIRouter(prefix='/api/annotations');guard=threading.RLock()

def now():return datetime.now(timezone.utc).isoformat()
def atomic(path,data):
    tmp=path.with_suffix('.tmp');tmp.write_text(json.dumps(data,ensure_ascii=False,allow_nan=False,indent=2));tmp.replace(path)
def project_path(project):
    try:
        if str(uuid.UUID(project))!=project:raise ValueError()
    except (ValueError,TypeError):raise HTTPException(404,'Проект не найден')
    p=PROJECTS/project
    if not (p/'project.json').exists():raise HTTPException(404,'Проект не найден')
    return p

def load_item(p,index):
    if index<0 or not (p/'items'/f'{index}.json').exists():raise HTTPException(404,'Изображение не найдено')
    return json.loads((p/'items'/f'{index}.json').read_text())

def view_image(path):
    """Unknown scanners may be annotated, but receive no unsupported clinical prediction."""
    try:
        a,m=read_dicom(path)
        return a,m,True
    except Exception:pass
    try:
        ds=pydicom.dcmread(path)
        if int(ds.get('NumberOfFrames',1))!=1 or int(ds.get('Rows',0))*int(ds.get('Columns',0))>16_000_000:raise ValueError('Размер / число кадров не поддерживается')
        if ds.get('PhotometricInterpretation') not in ('MONOCHROME1','MONOCHROME2'):raise ValueError('Нужен монохромный DICOM')
        a=ds.pixel_array.astype(float)
        if a.ndim!=2 or min(a.shape)<32 or not np.isfinite(a).all():raise ValueError('Некорректные пиксели')
        a=(255*(a-a.min())/max(float(np.ptp(a)),1e-9)).astype('uint8')
        if ds.PhotometricInterpretation=='MONOCHROME1':a=255-a
        spacing=ds.get('PixelSpacing');source='DICOM PixelSpacing' if spacing is not None else 'unknown'
        if spacing is not None:
            spacing=list(map(float,spacing))
            if len(spacing)!=2 or not np.isfinite(spacing).all() or min(spacing)<=0:spacing=None;source='unknown'
        return a,{'study_uid':str(ds.get('StudyInstanceUID','')),'image_uid':str(ds.get('SOPInstanceUID','')),'region':'unknown','side':'','spacing':spacing,'spacing_source':source},False
    except pydicom.errors.InvalidDicomError:pass
    with Image.open(path) as im:
        if im.width*im.height>16_000_000 or min(im.size)<32:raise ValueError('Размер изображения не поддерживается')
        a=np.array(ImageOps.exif_transpose(im).convert('L'))
    return a,{'study_uid':'','image_uid':'','region':'unknown','side':'','spacing':None,'spacing_source':'unknown'},False

def create_project(files,name,analyzer=None,base=None,progress=None):
    project=str(uuid.uuid4());p=PROJECTS/project
    for sub in ('items','images','originals'):(p/sub).mkdir(parents=True,exist_ok=True)
    info={'id':project,'name':str(name)[:160],'created_at':now(),'schema_version':1,'count':0,'errors':[],'model_proposals_are_expert_labels':False}
    compatible=[];mapping=[]
    try:
        for file in files:
            file=Path(file)
            try:
                if file.stat().st_size>32*1024**2:raise ValueError('Файл больше 32 MiB')
                a,m,supported=view_image(file);i=info['count'];original=f'{i}{file.suffix.lower() or ".dcm"}'
                shutil.copyfile(file,p/'originals'/original);Image.fromarray(a).save(p/'images'/f'{i}.png')
                digest=hashlib.sha256(file.read_bytes()).hexdigest();pixel_hash=hashlib.sha256(str(a.shape).encode()+a.tobytes()).hexdigest()
                rel=str(file.relative_to(base)) if base else file.name
                item={'index':i,'filename':rel,'original':original,'sha256':digest,'pixel_hash':pixel_hash,'study_uid':m['study_uid'],'image_uid':m['image_uid'],
                      'group_id':m['study_uid'] or (str(Path(rel).parent) if str(Path(rel).parent)!='.' else ''),
                      'width':int(a.shape[1]),'height':int(a.shape[0]),'region':m['region'],'side':m['side'],'spacing':m['spacing'],'spacing_source':m['spacing_source'],
                      'labels':{t:None for t in TASKS},'shapes':[],'status':'pending','revision':0,'reviewer':'','updated_at':now(),'proposal':None,
                      'geometry':None,'note':'','optional_conditions':{'scoliosis':None,'endoprosthesis':None}}
                item.update(projection='ap' if supported else 'unknown',field_of_view='full' if supported else 'unknown',
                            burned_roi='unknown',pixel_aspect_yx=None,hip_landmarks_visible=False,task_exclusions={})
                if supported:
                    item['geometry']=geometry(a,m)[1];compatible.append(file);mapping.append(i)
                atomic(p/'items'/f'{i}.json',item);info['count']+=1
            except Exception as e:info['errors'].append({'filename':file.name,'error':f'{type(e).__name__}: {e}'})
            if progress:progress(info['count'],len(files),'annotation_import')
        if not info['count']:raise ValueError('Не удалось открыть ни одного изображения')
        if analyzer and compatible:
            result=analyzer.analyze(compatible,p/'prelabels',base=base,progress=progress)
            for index,row,detail in zip(mapping,result['rows'],result['details']):
                item=load_item(p,index)
                if row['processing_status']=='Success':
                    probs=detail['probabilities'];thresholds=detail['thresholds']
                    item['proposal']={'model':detail['model'],'region':item['region'],'spacing':item['spacing'],'probabilities':{t:probs[t] for t in TASKS},'labels':{t:int(probs[t]>=thresholds[j+2]) for j,t in enumerate(TASKS)},'source':'machine_unconfirmed'}
                else:item['proposal_error']=detail.get('error','Inference failed')
                atomic(p/'items'/f'{index}.json',item)
        atomic(p/'project.json',info);return info
    except BaseException:
        shutil.rmtree(p,ignore_errors=True);raise

def validated(old,data,quick=False):
    item=dict(old)
    region=data.get('region',old['region'])
    if region not in ('spine','hip','other','unknown'):raise ValueError('Неизвестная область')
    side=data.get('side',old['side'])
    if side not in ('','left','right'):raise ValueError('Неизвестная сторона')
    spacing=data.get('spacing',old['spacing'])
    if spacing is not None:
        if not isinstance(spacing,list) or len(spacing)!=2 or any(isinstance(x,bool) or not isinstance(x,(int,float)) or not math.isfinite(x) or not 0<x<=100 for x in spacing):raise ValueError('Масштаб: два положительных числа [Y, X] мм/пиксель')
    labels=data.get('labels',old['labels'])
    if not isinstance(labels,dict) or set(labels)!=set(TASKS) or any(v is not None and (type(v) is not int or v not in (0,1)) for v in labels.values()):raise ValueError('Нужны пять меток: 0, 1 или null')
    labels=dict(labels);applicable=TASKS[:3] if region=='spine' else TASKS[3:] if region=='hip' else []
    for t in TASKS:
        if t not in applicable:labels[t]=None
    for key,allowed in [('projection',('unknown','ap','lateral','other')),('field_of_view',('unknown','full','cropped')),('burned_roi',('unknown','present','absent'))]:
        value=data.get(key,old.get(key,'unknown'))
        if value not in allowed:raise ValueError('Некорректное поле: '+key)
        item[key]=value
    ratio=data.get('pixel_aspect_yx',old.get('pixel_aspect_yx'))
    if ratio is not None and (type(ratio) not in (int,float) or not math.isfinite(ratio) or not .01<=ratio<=100):raise ValueError('Соотношение Y/X должно быть положительным числом от 0.01 до 100')
    visible=data.get('hip_landmarks_visible',old.get('hip_landmarks_visible',False))
    if type(visible) is not bool:raise ValueError('Видимость ориентиров должна быть логическим значением')
    exclusions=data.get('task_exclusions',old.get('task_exclusions',{}))
    if not isinstance(exclusions,dict) or any(t not in TASKS or not isinstance(reason,str) or not reason.strip() or len(reason)>500 for t,reason in exclusions.items()):raise ValueError('Для исключённого критерия нужна текстовая причина')
    item.update(region=region,spacing=spacing,pixel_aspect_yx=ratio,hip_landmarks_visible=visible,task_exclusions=dict(exclusions))
    blocks=task_blocks(item)
    if not quick and any(labels[t] is not None for t in applicable if t in blocks):raise ValueError('Недоступный критерий должен иметь метку null: '+', '.join(t for t in applicable if t in blocks and labels[t] is not None))
    state=data.get('status',old['status'])
    if state not in ('pending','confirmed','needs_review','skipped'):raise ValueError('Неизвестный статус')
    group=str(data.get('group_id',old['group_id'])).strip()[:256]
    if state=='confirmed':
        if quick:
            if region not in ('hip','spine') or any(labels[t] is None for t in applicable):raise ValueError('Отметьте нарушения или подтвердите норму')
        elif not applicable or any(labels[t] is None for t in applicable if t not in blocks) or not group:raise ValueError('Для подтверждения задайте область, группу исследования и все доступные метки; недоступные критерии остаются null')
    shapes=data.get('shapes',old['shapes'])
    if not isinstance(shapes,list) or len(shapes)>200:raise ValueError('Слишком много фигур')
    clean=[]
    for shape in shapes:
        kind=shape.get('type');points=shape.get('points');target=shape.get('target')
        if kind not in ('point','line','polyline','polygon','rect') or target not in TASKS+['axis','anatomy']:raise ValueError('Неизвестная фигура/задача')
        minpoints={'point':1,'line':2,'rect':2,'polyline':2,'polygon':3}[kind]
        if not isinstance(points,list) or not minpoints<=len(points)<=512 or (kind in ('point','line','rect') and len(points)!=minpoints):raise ValueError('Некорректное число точек')
        for point in points:
            if not isinstance(point,list) or len(point)!=2 or any(type(v) not in (int,float) or not math.isfinite(v) for v in point) or not 0<=point[0]<=old['width'] or not 0<=point[1]<=old['height']:raise ValueError('Координаты вне изображения')
        clean.append({'type':kind,'points':points,'target':target,'source':'human'})
    conditions=data.get('optional_conditions',old['optional_conditions'])
    if not isinstance(conditions,dict) or set(conditions)!={'scoliosis','endoprosthesis'} or any(v is not None and type(v) is not bool for v in conditions.values()):raise ValueError('Некорректные дополнительные признаки')
    item.update(region=region,side=side,spacing=spacing,labels=labels,shapes=clean,status=state,group_id=group,reviewer=str(data.get('reviewer',''))[:120],note=str(data.get('note',''))[:2000],optional_conditions=conditions,revision=old['revision']+1,updated_at=now())
    if spacing!=old['spacing']:item['spacing_source']='manual calibration' if spacing else 'unknown'
    # A proposal remains visible with its original provenance, but is flagged after metadata changes.
    item['proposal_stale']=bool(item['proposal'] and (region!=item['proposal']['region'] or spacing!=item['proposal']['spacing']))
    return item

def dataset_settings(p):
    info=json.loads((p/'project.json').read_text());first=load_item(p,0)
    key=info.get('dataset_key') or first.get('source_dataset') or info['id']
    # Queues split by anatomy still share their source dataset's calibration.
    token=hashlib.sha256(key.encode()).hexdigest()
    path=PROJECTS/'_dataset_settings'/f'{token}.json'
    if path.exists():return path,json.loads(path.read_text())
    spacing=first.get('spacing') or [1.05,.6]
    assumed=first.get('spacing') is None or first.get('spacing_assumed',False)
    return path,{'dataset_key':key,'revision':0,'spacing':spacing,'spacing_assumed':assumed,
                 'spacing_source':'dataset working default' if assumed else first.get('spacing_source','dataset default')}

@router.get('/{project}/settings')
def get_settings(project:str):return dataset_settings(project_path(project))[1]

@router.put('/{project}/settings')
async def save_settings(project:str,request:Request):
    try:data=await request.json()
    except ValueError:raise HTTPException(422,'Некорректный JSON')
    if not isinstance(data,dict):raise HTTPException(422,'Нужен объект JSON')
    with guard:
        path,old=dataset_settings(project_path(project))
        if data.get('revision')!=old['revision']:raise HTTPException(409,'Настройки датасета изменены. Перезагрузите страницу.')
        spacing=data.get('spacing')
        if not isinstance(spacing,list) or len(spacing)!=2 or any(type(x) not in (float,int) or not math.isfinite(x) or not 0<x<=100 for x in spacing):raise HTTPException(422,'Нужны положительные Y и X в мм/пиксель')
        result=old|{'spacing':spacing,'revision':old['revision']+1,'updated_at':now()}
        if spacing!=old['spacing']:result.update(spacing_assumed=False,spacing_source='dataset user-specified scale')
        path.parent.mkdir(parents=True,exist_ok=True);atomic(path,result);return result

def quick_region(row):
    proposal=row.get('anatomy_proposal') or {}
    if row.get('projection') in ('lateral','other') or proposal.get('projection') in ('lateral','other'):return 'excluded'
    region=row['region'] if row['region'] in ('spine','hip') else proposal.get('region','unknown')
    return region if region in ('spine','hip') else 'unknown'

@router.get('')
def list_projects():
    PROJECTS.mkdir(parents=True,exist_ok=True);out=[]
    for f in PROJECTS.glob('*/project.json'):
        try:out.append(json.loads(f.read_text()))
        except (ValueError,OSError):continue
    return sorted(out,key=lambda p:p['created_at'],reverse=True)

@router.get('/{project}')
def get_project(project:str):
    p=project_path(project);info=json.loads((p/'project.json').read_text());items=[]
    for i in range(info['count']):
        r=load_item(p,i);probs=(r['proposal'] or {}).get('probabilities',{});tasks=TASKS[:3] if r['region']=='spine' else TASKS[3:]
        uncertainty=max((1-abs(2*probs[t]-1) for t in tasks if t in probs),default=1.)
        items.append({k:r[k] for k in ('index','filename','region','status','revision')}|{'uncertainty':uncertainty,'quick_region':quick_region(r)})
    return info|{'items':items,'settings':dataset_settings(p)[1]}

@router.get('/{project}/items/{index}')
def get_item(project:str,index:int,quick:bool=False):
    p=project_path(project);row=load_item(p,index)
    if quick and quick_region(row)=='spine':
        settings=dataset_settings(p)[1];a=np.array(Image.open(p/'images'/f'{index}.png'))
        row['display_geometry']=geometry(a,{'region':'spine','spacing':settings['spacing'],'spacing_source':settings['spacing_source']})[1]
    return row

@router.put('/{project}/items/{index}')
async def save_item(project:str,index:int,request:Request):
    payload=await request.body()
    if len(payload)>2*1024**2:raise HTTPException(413,'Разметка слишком большая')
    try:data=json.loads(payload)
    except ValueError:raise HTTPException(422,'Некорректный JSON')
    if not isinstance(data,dict):raise HTTPException(422,'Нужен объект JSON')
    with guard:
        p=project_path(project);old=load_item(p,index)
        if data.get('revision')!=old['revision']:raise HTTPException(409,'Разметка изменена другим оператором. Перезагрузите снимок перед сохранением.')
        quick=data.get('review_mode')=='quick'
        if quick:
            settings=dataset_settings(p)[1]
            if data.get('settings_revision')!=settings['revision']:raise HTTPException(409,'Масштаб датасета изменился. Перезагрузите страницу перед подтверждением.')
            if data.get('region') not in ('spine','hip'):raise HTTPException(422,'Выберите бедро или позвоночник')
            # Preserve acquisition provenance; a quick review judges the displayed image.
            data={**data,'projection':'ap','spacing':settings['spacing'],'field_of_view':old.get('field_of_view','unknown'),
                  'hip_landmarks_visible':old.get('hip_landmarks_visible',False),'task_exclusions':{},'pixel_aspect_yx':None}
        try:item=validated(old,data,quick=quick)
        except (ValueError,TypeError,AttributeError) as e:raise HTTPException(422,str(e))
        if quick:
            item.update(annotation_scope='displayed_image',review_mode='quick',settings_revision=settings['revision'],dataset_key=settings['dataset_key'],
                        spacing_source=settings['spacing_source'],spacing_assumed=settings['spacing_assumed'])
        if item['region']=='spine' and item['spacing'] is not None:
            a=np.array(Image.open(p/'images'/f'{index}.png'));item['geometry']=geometry(a,{'region':'spine','spacing':item['spacing'],'spacing_source':item['spacing_source']})[1]
        else:item['geometry']=None
        # Keep full revisions so corrections and reviewer provenance remain recoverable.
        history=p/'history'/str(index);history.mkdir(parents=True,exist_ok=True)
        atomic(history/f"{old['revision']}.json",old)
        atomic(p/'items'/f'{index}.json',item)
        return item

@router.get('/{project}/images/{index}.png')
def preview(project:str,index:int):
    p=project_path(project);load_item(p,index);return FileResponse(p/'images'/f'{index}.png')

def export_records(p,confirmed_only):
    info=json.loads((p/'project.json').read_text());out=[]
    for i in range(info['count']):
        row=load_item(p,i)
        if confirmed_only and row['status']!='confirmed':continue
        row.update(export_admission(row))
        row['project_id']=info['id'];row['original_relative_path']=f"originals/{row['original']}";out.append(row)
    return out

@router.get('/{project}/export')
def export(project:str,format:str='jsonl',confirmed_only:bool=True):
    import io
    p=project_path(project)
    with guard:rows=export_records(p,confirmed_only)
    if format=='jsonl':return Response(''.join(json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n' for r in rows),media_type='application/x-ndjson',headers={'Content-Disposition':'attachment; filename="annotations.jsonl"'})
    if format=='xlsx':
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.title='Annotations';columns=['index','filename','study_uid','group_id','region','projection','field_of_view','burned_roi','side','status','reviewer','revision','quality_class',*TASKS,'annotation_scope','review_mode','settings_revision','spacing_assumed','training_mask','task_blocks','pixel_aspect_yx','hip_landmarks_visible','source_dataset','source_license','source_training_allowed','spacing','spacing_source','shapes','sha256','pixel_hash','note'];ws.append(columns)
        for row in rows:
            flat=row|row['labels'];ws.append([json.dumps(flat.get(k),ensure_ascii=False) if isinstance(flat.get(k),(dict,list)) else flat.get(k) for k in columns])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;b=io.BytesIO();wb.save(b)
        return Response(b.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="annotations.xlsx"'})
    raise HTTPException(400,'format=jsonl или xlsx')
